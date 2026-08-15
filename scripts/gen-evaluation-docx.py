#!/usr/bin/env python3
"""
Build the branch-evaluation DOCX for integration/main.

Companion to scripts/gen-arch-docx.py and deliberately in the same visual
language, but a different document: that one describes what the system IS,
this one judges whether it WORKS and what it would take to run it over a
14,000-bus fleet.

Every figure quoted is either measured in the repository (and cited to the
file that measured it) or a stated arithmetic projection. Verification
results are from an actual run, not a claim.

Run:  python3 scripts/gen-evaluation-figures.py && python3 scripts/gen-evaluation-docx.py
Out:  docs/branch-evaluation/Olympuss_AI_Branch_Evaluation_integration-main.docx
"""
import os
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIGURES = os.path.join(ROOT, "docs", "branch-evaluation", "figures")
OUT = os.path.join(ROOT, "docs", "branch-evaluation",
                   "Olympuss_AI_Branch_Evaluation_integration-main.docx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

INK = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
MUTED = RGBColor(0x47, 0x55, 0x69)
RED = RGBColor(0xB9, 0x1C, 0x1C)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREEN = RGBColor(0x15, 0x6F, 0x3C)

doc = Document()

# ---- base styles ----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for lvl, sz, col in [("Heading 1", 17, BLUE), ("Heading 2", 13.5, INK), ("Heading 3", 11.5, TEAL)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = col
    st.font.bold = True

mono = doc.styles.add_style("Mono", WD_STYLE_TYPE.CHARACTER)
mono.font.name = "Consolas"
mono.font.size = Pt(9.5)
mono.font.color.rgb = RGBColor(0x0B, 0x3D, 0x2E)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def code(paragraph, text):
    r = paragraph.add_run(text)
    r.style = doc.styles["Mono"]
    return r


def para(text="", size=None, bold=False, italic=False, color=None, after=6, before=0, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def rich(segments, after=6, before=0, style=None):
    """segments: list of (text, kind) where kind in {'', 'b', 'i', 'code'}."""
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    for text, kind in segments:
        if kind == "code":
            code(p, text)
        else:
            r = p.add_run(text)
            r.bold = kind == "b"
            r.italic = kind == "i"
    return p


def bullet(segments):
    return rich(segments, after=3, style="List Bullet")


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def h3(text):
    doc.add_heading(text, level=3)


def callout(label, text, fill="FEF3C7", label_color=AMBER):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.color.rgb = label_color
    r.font.size = Pt(9.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def table(headers, rows, widths=None, header_fill="1D4ED8", fs=8.6, caption=None):
    if caption:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_after = Pt(2)
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_bg(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]
        rr = p.add_run(htext)
        rr.bold = True
        rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        rr.font.size = Pt(fs)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            for j, seg in enumerate(val.split("`")):
                if seg == "":
                    continue
                if j % 2 == 1:
                    code(p, seg)
                else:
                    rr = p.add_run(seg)
                    rr.font.size = Pt(fs)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(filename, caption, width=6.5):
    path = os.path.join(FIGURES, filename)
    if not os.path.exists(path):
        para(f"[missing figure: {filename} — run scripts/gen-evaluation-figures.py]", color=RED)
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED


def page_break():
    doc.add_page_break()


def add_page_number(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)


# =====================================================================================
# COVER
# =====================================================================================
for _ in range(3):
    doc.add_paragraph()
para("OLYMPUSS AI", size=13, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("UPSRTC — Uttar Pradesh State Road Transport Corporation",
     size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
para("Branch Evaluation\nintegration/main", size=28, bold=True, color=INK,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
para("What was built · whether it works · what it would take to detect and\n"
     "resolve bus bunching across a fleet of 14,000 buses",
     size=11.5, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
para("An independent read of all 98 commits on the branch, verified by running the "
     "test suites, typecheck and lint in both packages.",
     size=10.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para(f"98 commits · 738 files · +140,217 / −18,494 lines against `main`".replace("`", ""),
     size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(f"Prepared {date.today().strftime('%d %B %Y')}", size=9.5, color=MUTED,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
para("Every quantity in this document is either measured in the repository and cited to "
     "the file that measured it, or a stated arithmetic projection labelled as such. "
     "No secret values are printed.",
     size=8.5, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# =====================================================================================
# TOC
# =====================================================================================
h1("Contents")
para("In Microsoft Word, right-click the field below and choose “Update Field” "
     "(or select it and press F9) to populate the entries and page numbers.",
     size=9, italic=True, color=MUTED)
toc_p = doc.add_paragraph()
run = toc_p.add_run()
b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve")
i.text = r'TOC \o "1-2" \h \z \u'
s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
t_ = OxmlElement("w:t"); t_.text = "Right-click and Update Field to build the table of contents."
e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
for el in (b, i, s, t_, e):
    run._r.append(el)
page_break()

# ---- page setup, header, footer ----
section = doc.sections[0]
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.8)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Olympuss AI / UPSRTC — Branch Evaluation: integration/main   ·   Page ")
add_page_number(fp)
for r in fp.runs:
    r.font.size = Pt(8); r.font.color.rgb = MUTED
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("OLYMPUSS AI · UPSRTC · CONFIDENTIAL ENGINEERING DOCUMENT")
hr.font.size = Pt(7.5); hr.font.color.rgb = MUTED; hr.bold = True

# =====================================================================================
# 1. EXECUTIVE SUMMARY
# =====================================================================================
h1("1. Executive Summary")

para("`main` is a UI prototype. `integration/main` is the actual system: two independently "
     "deployed applications, a real GPS ingestion and state-estimation pipeline, a headway "
     "controller, a human-authorised command path that reaches a driver's phone, and a "
     "role-based operations console on top of all of it.".replace("`", ""))

callout("VERDICT",
        "The engineering quality is high and the honesty discipline is exceptional. Everything "
        "builds, typechecks, lints and passes 2,410 tests. It is a working, end-to-end system at "
        "pilot scale on a handful of corridors. It is not yet a 14,000-bus system, for two "
        "independent reasons — throughput and coverage — plus one architectural gap that changes "
        "what the product is at fleet scale.",
        fill="DBEAFE", label_color=BLUE)

h3("Scorecard")
table(
    ["Dimension", "Verdict", "Basis"],
    [
        ["Does it build and pass its tests",
         "YES — fully green",
         "2,410 tests pass; typecheck and lint clean in both packages (§3)"],
        ["Is the bunching loop closed end to end",
         "YES — detect → propose → approve → deliver → acknowledge",
         "Verified by reading every stage plus two E2E specs (§2)"],
        ["Is it honest about what it measures",
         "YES — unusually so; a genuine strength",
         "Calibration refuses to invent H*; consoles refuse to print unobserved numbers (§4)"],
        ["Can it see the fleet",
         "NO — 7.2% of vehicles report live GPS",
         "664 of 9,262 records, `harvest.ts` (§5.1)"],
        ["Can it detect across the network",
         "NO — 25.7% of route-directions have a target headway",
         "171 of 666, `odTimetable.ts` (§5.2)"],
        ["Does it propose actions network-wide",
         "NO — the solver is pull-only, one corridor at a time",
         "`solve()` has exactly one caller, `POST /v1/mpc/solve` (§5.3)"],
        ["Will ingestion sustain 14,000 buses",
         "NO — sequential loop, ~84–210 s per 30 s cycle",
         "Projection from the query count in `pipeline.ts` / `service.ts` (§6.1)"],
        ["Will the map render 14,000 buses",
         "YES — already benchmarked past it",
         "9,170 vehicles at 4.3 ms/frame, `fleetCanvasLayer.ts` (§6.5)"],
    ],
    widths=[2.0, 2.1, 2.9])

h3("The three things to fix first")
rich([("1.  Batch the ingestion writes.", "b"),
      ("  The single highest-value change in the repository. It converts a projected "
       "84–210 second ingestion cycle into single-digit seconds and is the difference "
       "between a controller that acts and one that silently stops acting (§6.1).", "")], after=4)
rich([("2.  Add a scheduled network-wide solve and a ranked recommendation queue.", "b"),
      ("  Today an operator must select each corridor to discover it has a problem. "
       "This is the difference between a tool that works and a tool that gets used (§5.3).", "")], after=4)
rich([("3.  Raise GPS reporting above 7.2%.", "b"),
      ("  Not a code change — a procurement and operations conversation with UPSRTC. "
       "Nothing built downstream matters more than this (§5.1).", "")], after=4)

page_break()

# =====================================================================================
# 2. WHAT LANDED
# =====================================================================================
h1("2. What Landed on the Branch")

h2("2.1  The size of the change")
table(
    ["Area", "Files", "Added", "Removed", "What it is"],
    [
        ["`control-service/src`", "111", "+20,315", "0", "The new backend service — did not exist on `main`"],
        ["`control-service/test(s)`", "79", "+23,121", "0", "Its unit and integration suites"],
        ["`src/components`", "137", "+16,300", "−3,176", "Operator consoles and the driver surface, rebuilt"],
        ["`src/tests`", "113", "+29,171", "−656", "Web-app suites"],
        ["`src/lib`", "102", "+14,371", "−2,298", "Auth/RBAC, control-service clients, ops models, copilot"],
        ["`src/app`", "84", "+8,173", "−385", "Routes, pages and API handlers"],
        ["`db/migrations` + `control-service/db`", "19", "+2,619", "0", "Two separate datastores, by design"],
        ["`tests/e2e`", "9", "+1,608", "−59", "Playwright specs incl. live command delivery"],
    ],
    widths=[1.9, 0.5, 0.75, 0.8, 3.05],
    caption="Grouped from `git diff --numstat main..integration/main`. Totals: 738 files, +140,217 / −18,494.")

h2("2.2  The six workstreams")
table(
    ["Workstream", "What landed"],
    [
        ["A second deployed service",
         "`control-service/` — Express + Postgres/PostGIS with its own migrations, deployed separately from the "
         "Next.js app. Ingestion, state estimation, headway, MPC, commands, pilot rollout, rehearsal, arrival prediction."],
        ["Sensing",
         "In-process GPS poller → map-matching (Kalman filter + a 0.05° spatial grid prefilter) → `vehicle_states`. "
         "A network-geometry cache with single-flight loading and stale-while-revalidate refresh removes what its own "
         "header calls “the wall that has to come down before ingestion can run at fleet scale”."],
        ["Detection",
         "Time-domain headway, EWT and CV per route-direction; a reactive rule that opens a `bunching_incident` only "
         "after k consecutive samples below a configured fraction of H*, and closes it on recovery."],
        ["Decision and delivery",
         "Terminal-dispatch regulation, two-way holding, self-equalizing fallback, a hard safety filter, and an "
         "occupancy-weighted advisory — then dispatcher approval, `commands`, webhook delivery, and a driver PWA "
         "that acknowledges."],
        ["Identity and consoles",
         "Supabase-backed RBAC across seven roles with depot scoping, invites, an append-only audit log, kill switches "
         "and rollout stages; driver, dispatcher, depot, control-room, planner and admin consoles rebuilt on shadcn/ui "
         "+ Tailwind."],
        ["Calibration and provenance",
         "Target headway derived from two published UPSRTC sources with a precedence rule, never fabricated. Every "
         "screen distinguishes measured data from model output (`docs/LIVE_VS_PREDICTED.md` is authoritative)."],
    ],
    widths=[1.55, 5.45])

h2("2.3  Not committed")
bullet([("`docs/technical-architecture/` — a generated architecture .docx plus 13 diagrams, and the two "
         "Python generators `scripts/gen-arch-diagrams.py` and `scripts/gen-arch-docx.py`. These are untracked "
         "in the working tree. They are build outputs and tooling, not product code, but the generators are "
         "worth committing so the document can be rebuilt by anyone.", "")])

page_break()

# =====================================================================================
# 3. VERIFICATION
# =====================================================================================
h1("3. Verification — Does It Actually Work?")

para("Not taken on trust. Both packages were installed and everything below was run.")

table(
    ["Check", "Command", "Result"],
    [
        ["control-service tests", "`pnpm test` in `control-service/`", "731 passed, 63 files, 3.3 s"],
        ["Web-app tests", "`pnpm test` at the repo root", "1,679 passed, 46 skipped, 102 files"],
        ["Typecheck — web", "`pnpm typecheck`", "Clean"],
        ["Typecheck — control-service", "`pnpm typecheck`", "Clean"],
        ["Lint — web", "`pnpm lint`", "No warnings or errors"],
        ["Lint — control-service", "`pnpm lint`", "Clean"],
    ],
    widths=[1.7, 2.5, 2.8],
    caption="2,410 tests pass in total. Executed locally against the branch as checked out.")

h3("On the 46 skipped tests")
para("Five test files skip locally and are designed to. They are the ones that prove guarantees living in SQL "
     "rather than in TypeScript — keyset pagination that must not lose safety records, `ON CONFLICT` dedupe — "
     "so they need a real `OPS_DATABASE_URL`. They skip when it is unset and hard-FAIL when `CI=true`. "
     "The CI workflow provisions an `ops-db` service and runs `pnpm migrate:ops` against it, so nothing skips in CI. "
     "That guard exists because those files skipped on every CI run the repository had ever done — a green suite "
     "that said nothing about the SQL it was supposed to protect.")

h3("One environment issue found, and it was not a code defect")
para("Two test files initially failed to load: `@radix-ui/react-slot` could not be resolved from "
     "`src/components/ui/button.tsx`. The package is declared in `package.json` and present in `pnpm-lock.yaml`; "
     "the local `node_modules` was simply stale, predating the commit that introduced shadcn/ui. Running "
     "`pnpm install --frozen-lockfile` installed it — and pruned five other `@radix-ui/*` packages that nothing "
     "in `src` imports — after which the full suite passed. Worth noting only because a newcomer cloning the "
     "branch will hit the same thing if they skip the install step in either package.")

h3("What was not verified, and why")
bullet([("The Playwright E2E specs", "b"),
        (" were not run. They require a live control-service, a running web app, a seeded Postgres and seven "
         "seeded operator accounts. CI provisions all of it and runs them; that is the right place for them.", "")])
bullet([("Behaviour against the live UPSRTC feed", "b"),
        (" was not exercised. No credentials were used and no upstream endpoint was contacted.", "")])
bullet([("The throughput numbers in §6 are projections, not benchmarks.", "b"),
        (" They are computed from the query count the ingestion path demonstrably issues. The absence of any "
         "load test in the repository is itself one of the findings.", "")])

page_break()

# =====================================================================================
# 4. WHAT IS GOOD
# =====================================================================================
h1("4. What Is Genuinely Good")

para("Three things stand out, and they are the reason the gaps in §5 and §6 are worth fixing rather than "
     "working around.")

h2("4.1  It refuses to invent numbers, at a measured cost")
para("The clearest example is the OD calibration seeder. Two estimators were run over the same 53,898-row corpus "
     "and compared on the 906 line-directions both could answer for:")
table(
    ["Estimator", "Median derived H*", "Line-directions answered"],
    [["Pooled per route", "1,635 s", "1,441"],
     ["Bucketed per boarding stop", "2,582 s", "940"]],
    widths=[2.4, 2.0, 2.6],
    caption="Measured over the full sweep of all 210 ordered city pairs, 2026-08-11 (`odTimetable.ts`).")
para("Pooling reports a 37% shorter headway on 44% of the pairs. Because every threshold in the detection path "
     "is a ratio of H*, a too-small denominator makes a bunched route look healthy — the silent false negative. "
     "Against the seeded database, pooling would have calibrated 111 route-directions and bucketing calibrates 97. "
     "Fourteen route-directions of coverage were deliberately given up to keep the other 97 honest, and the trade "
     "is documented numerically rather than asserted.")

h2("4.2  The failure modes are the right ones")
bullet([("The rollout gate fails closed.", "b"),
        (" It previously returned early and ALLOWED a command when an approval carried no route-direction — a "
         "fail-open hole with the blast radius of the entire gate. Both paths now return 422.", "")])
bullet([("Arrival prediction declines rather than extrapolating.", "b"),
        (" Its ten-minute cutoff is derived, not chosen: at the measured 38.5 km/h median running speed a bus "
         "covers 6.4 km in ten minutes, against a measured 7.9 km median inter-stop gap.", "")])
bullet([("The copilot returns 503 rather than a guess.", "b"),
        (" There is deliberately no fallback path that produces text without a successful model call.", "")])
bullet([("The map shows an outage as an outage.", "b"),
        (" Upstream failure yields zero rows and an `UPSTREAM UNAVAILABLE` badge — never substituted demo buses "
         "in the same table as real ones.", "")])

h2("4.3  The comments encode incidents, not intentions")
para("This is rarer than it sounds and it is what will let a new engineer work safely in this code.")
bullet([("`headwayCompute.ts` explains why computing headway on the ingestion path is catastrophic rather than "
         "merely expensive: it would collapse three samples taken three minutes apart into 45 milliseconds, "
         "turning every GPS wobble into an incident. “The cadence IS the semantics.”", "")])
bullet([("`store.ts#upsertHeadwayStates` records the bug it exists to fix — the solver read a permanently empty "
         "headway map and returned zero candidates on a genuinely bunched route. “Persisting is not publishing.”", "")])
bullet([("The command-lifecycle notes record a delivery path that told an operator an approval had not been "
         "consumed while the driver already had the instruction on screen, and why one webhook dispatch is "
         "deliberately not awaited as a result.", "")])

page_break()

# =====================================================================================
# 5. EFFECTIVENESS GAPS
# =====================================================================================
h1("5. Effectiveness Gaps")

para("These are not defects. Everything here works as designed; the question is how much of the bunching problem "
     "the design can currently reach.")

figure("01-coverage-ceilings.png",
       "Figure 1 — The two coverage ceilings. Both are upstream of anything the control logic does, and neither "
       "is fixed by writing code.")

h2("5.1  Only 7.2% of the fleet is visible")
para("Measured from the upstream feed on 2026-08-09 (`harvest.ts`): 9,262 vehicle records, of which 664 carry "
     "`status === 'Live'`. Only 1,602 carry a route name at all. Among the live subset there are 517 distinct "
     "route names.")
callout("CONSEQUENCE",
        "Nothing downstream can compensate for this. Headway is a property of a pair of buses; a bus that does "
        "not report a position cannot be half of one.",
        fill="FEE2E2", label_color=RED)

h2("5.2  Only 25.7% of the network can detect at all")
para("171 of 666 active route-directions have a target headway — 74 from the corridor timetable and 97 from the "
     "statewide OD schedule. The remaining 495 have none, so detection is structurally off for them. Correctly and "
     "visibly off, but off.")
para("The seeder names its own ceiling precisely: `line_name` matches `routes.id` for only about 170 of the 666 "
     "seeded route-directions, because `routes.id` came from one feed's `line_id` while this feed's `line_name` is "
     "frequently a description rather than an identifier. The limit is the join, not the estimator — which means it "
     "is fixable without touching any control logic.")
para("Compounding it: with 664 live buses spread across 517 route names, most route-directions carry zero or one "
     "vehicle and cannot form a leader/follower pair at all. The sweep's own header puts the number that can at "
     "“a few dozen”.")

h2("5.3  Nothing proposes an action unless a human goes looking")
para("This is the most consequential architectural gap in the branch.")
para("`solve()` has exactly one caller in the entire repository: `POST /v1/mpc/solve`, which the web app invokes "
     "when an operator opens the control room on one selected corridor. The service runs five scheduled jobs — "
     "command delivery sweep, TTL sweep, GPS poll, headway compute and geometry refresh. None of them is the solver.")
table(
    ["Stage", "Runs automatically, network-wide?", "Mechanism"],
    [["Detection", "Yes", "`headwayCompute` sweep every 60 s"],
     ["Resolution", "No", "Operator opens a corridor → one `POST /v1/mpc/solve`"]],
    widths=[1.3, 2.3, 3.4])
para("The human-approval requirement is correct and should stay — it is named an explicit permanent design "
     "constraint. The gap is not the approval. It is that nothing proposes across the network and ranks by "
     "severity, so an operator responsible for 666 route-directions has to guess which one to look at.")

h2("5.4  The control authority does not match the network")
para("Easy to miss, and it undercuts the whole control layer. The seeded network averages 215 km and 19 stops per "
     "route-direction, with a measured 7.9 km median gap between stops and a median OD-derived target headway of "
     "2,582 seconds — roughly 43 minutes. `max_hold_seconds` defaults to 90.")
callout("THE MISMATCH",
        "Classic bunching control is an urban high-frequency problem: 5–15 minute headways, where a 60-second hold "
        "is meaningful authority. Against a 43-minute headway a 90-second hold corrects 3.5% of the target gap. "
        "The control laws are correctly implemented; they are simply under-powered for this operating regime.",
        fill="FEF3C7", label_color=AMBER)

h2("5.5  It is not actually model-predictive control")
para("`prediction_horizon_control_points` is used, in the code's own words, “to label/scale the occupancy-weighted "
     "MPC advisory, not to run a true multi-step solve yet”. The three control laws are reactive proportional "
     "controllers over the current headway snapshot. There is no receding-horizon optimisation and no forward "
     "prediction of where a bunch travels. The code is honest about this; the naming is not.")

h2("5.6  The data the model needs does not exist yet")
para("From the rehearsal simulator's provenance manifest: `route_links` is empty, `trips` and `trip_stop_times` "
     "are empty, and `vehicle_states.occupancy_count` is null fleet-wide.")
table(
    ["Missing input", "What it disables"],
    [["Occupancy counts", "The occupancy-weighted advisory always runs on an estimated mid-load — the weighting is inert"],
     ["`trips` / `trip_stop_times`", "`resolveCurrentTrip` always returns null; schedule adherence is unmeasurable"],
     ["`route_links`", "No segment-level running-time model"]],
    widths=[1.9, 5.1])
para("The rehearsal simulator is scrupulous about this: its control laws are the deployed ones, and its demand, "
     "dwell, capacity and disturbance model is openly invented. That line — real control, guessed traffic — is "
     "stated field by field in every rehearsal's provenance manifest.")

h2("5.7  Three levers, all of them holds")
para("The solver emits `terminal_dispatch_hold`, `two_way_hold` and `self_equalizing_hold`. The schema already "
     "models `stop_skip`, `short_turn`, `deadhead`, `boarding_limit` and `standby_injection`; none is generated. "
     "On 215 km routes, short-turn and standby injection are the levers with real authority. Holding is the weakest "
     "one available, and it is the only one implemented.")

page_break()

# =====================================================================================
# 6. SCALE
# =====================================================================================
h1("6. Scaling to 14,000 Buses")

para("At today's 664 live buses every component fits comfortably. What follows is what breaks at roughly "
     "twenty-one times that, in the order it breaks.")

h2("6.1  Wall 1 — ingestion is sequential (hard blocker)")
para("`ingestPositionEvents` is a `for` loop with `await` inside it. Each event triggers three database "
     "round-trips: `loadActiveHold`, `resolveCurrentTrip` and `saveVehicleState`. The batch therefore costs "
     "fleet × 3 × round-trip, strictly serially.")

figure("02-ingestion-throughput.png",
       "Figure 2 — One ingestion cycle against the 30-second poll interval it must fit inside. Today's 664 live "
       "buses sit near the origin, comfortably within budget.")

callout("THE FAILURE IS QUIET, WHICH IS WHAT MAKES IT DANGEROUS",
        "Overlap suppression means slow cycles are skipped rather than stacked, so position age drifts to 1.5–3.5 "
        "minutes. The hard safety filter then rejects every candidate as stale against its 90-second bound. The "
        "system fails safe and stops acting — while the logs still show healthy polls. Nothing looks broken.",
        fill="FEE2E2", label_color=RED)

h3("The fix")
bullet([("Batch the writes.", "b"), ("  A multi-row `INSERT … ON CONFLICT` carrying the same "
         "`observed_at` out-of-order guard, driven by `unnest`.", "")])
bullet([("Prefetch the holds.", "b"), ("  One query for the whole batch instead of one per vehicle.", "")])
bullet([("Drop `resolveCurrentTrip` until `trips` is populated.", "b"),
        ("  It always returns null today (§5.6), so it is a round-trip per bus per cycle for a guaranteed null.", "")])

h2("6.2  Wall 2 — detection latency multiplies")
para("The sweep visits a route-direction once every ⌈eligible ÷ batch⌉ cycles, so latency does not degrade "
     "smoothly — it jumps by a whole three-minute sample window at a time.")

figure("03-detection-latency.png",
       "Figure 3 — Time to first detection as the eligible set outgrows one sweep. The code already logs a warning "
       "when this starts happening; it does not mitigate it.")

para("Raising `HEADWAY_BATCH_SIZE` is the obvious fix and immediately runs into "
     "`HEADWAY_COMPUTE_CONCURRENCY = 4` against a connection pool capped at 10. All three have to move together, "
     "and the pool has to move first.")

h2("6.3  Wall 3 — single process by construction")
table(
    ["Component", "Why it cannot scale horizontally today"],
    [["`stateStore`", "A module-level in-memory `Map`, rehydrated from Postgres at boot"],
     ["Network geometry cache", "Per-process, TTL'd and self-refreshing"],
     ["GPS poller", "`GPS_POLL_ENABLED` must be true on exactly one instance, or replicas duplicate ingestion and race the out-of-order guard"],
     ["`listVehicleStates(routeDirectionId)`", "Full O(N) scan of every vehicle, per solve"]],
    widths=[2.1, 4.9])
para("The whole ingest-to-detect path is therefore single-node. Moving the state store behind Redis is already "
     "named in the production roadmap as Phase 4 work; at 14,000 buses it stops being hardening and becomes a "
     "prerequisite.")

h2("6.4  Wall 4 — the copilot spawns a CLI per request")
para("The copilot authenticates with an operator's personal Claude subscription OAuth token and invokes the "
     "`claude` CLI as a child process per call. The reasoning is sound for the current setup and the module is "
     "careful — it strips `ANTHROPIC_API_KEY` from the child environment so a stray key cannot silently move "
     "calls onto metered billing. But a control room with concurrent users will hit subscription rate limits, pay "
     "a process spawn per request, and require the CLI installed on every server instance.")

h2("6.5  What already scales")
para("The fleet map. Measured at 9,170 vehicles and 4.3 ms per frame — 120 fps sustained, down from 53.7 ms per "
     "frame and 18.5 fps before the rewrite. It will render 14,000 buses without modification.")

page_break()

# =====================================================================================
# 7. REMAINING WORK
# =====================================================================================
h1("7. What Remains, Prioritised")

h2("7.1  Blocking for 14,000 buses")
table(
    ["#", "Work", "Why it is first"],
    [["1", "Batch ingestion writes and the hold lookup",
      "Converts a projected 84–210 s cycle into single digits. Highest value-per-line change in the repository."],
     ["2", "Raise the pg pool; re-tune `HEADWAY_BATCH_SIZE` and `HEADWAY_COMPUTE_CONCURRENCY` together; "
           "add a 14,000-fixture load test",
      "There is no throughput budget stated anywhere in the repository and no load test. Both are gaps in their own right."],
     ["3", "Move `stateStore` behind Redis so ingestion can shard by route-direction",
      "The only route to more than one ingesting instance."]],
    widths=[0.35, 2.85, 3.8])

h2("7.2  Blocking for effectiveness")
table(
    ["#", "Work", "Why"],
    [["4", "A scheduled network-wide solve feeding a severity-ranked recommendation queue",
      "So the control room opens on “here are the twelve corridors that need you”, not a corridor picker. "
      "The difference between a tool that works and one that gets used."],
     ["5", "Fix the calibration join to lift H* coverage past 26%",
      "The blocker is identifier-space mismatch between two upstream feeds, not the estimator — so the fix is "
      "cheap relative to its effect."],
     ["6", "Raise GPS reporting above 7.2% with UPSRTC",
      "Not a code change. Nothing built downstream matters more."]],
    widths=[0.35, 2.85, 3.8])

h2("7.3  Effectiveness, next tier")
table(
    ["#", "Work", "Why"],
    [["7", "Per-corridor `max_hold_seconds` matched to a 43-minute H*, or implement `short_turn` / `standby_injection`",
      "Holding is the weakest lever and currently the only one. The schema already models the others."],
     ["8", "Populate `trips` and `trip_stop_times`",
      "Unlocks schedule adherence and makes `resolveCurrentTrip` meaningful instead of a guaranteed null."],
     ["9", "Real receding-horizon MPC over `prediction_horizon_control_points`",
      "Replaces one-step reactive control and makes the module's name true."],
     ["10", "Ingest occupancy from ticketing or APC",
      "Until then the occupancy-weighted advisory is decorative."]],
    widths=[0.35, 2.85, 3.8])

h2("7.4  Hardening")
table(
    ["#", "Work", "Why"],
    [["11", "Replace the copilot CLI spawn with a proper API integration",
      "Before multi-user load, not after."],
     ["12", "SSE or WebSocket push instead of polling",
      "Already on the roadmap; matters more as corridor count grows."],
     ["13", "Refresh `docs/PRODUCTION_ROADMAP.md`",
      "It still lists “No authentication” as a known gap. This branch closed it; the doc is partly stale and will "
      "mislead a reader who trusts it."]],
    widths=[0.35, 2.85, 3.8])

page_break()

# =====================================================================================
# APPENDIX A
# =====================================================================================
h1("Appendix A — Measured Figures and Their Sources")

para("Every quantity used in this document, with the file that measured it. Nothing here is estimated unless the "
     "row says so.")

table(
    ["Quantity", "Value", "Source"],
    [
        ["Vehicle records in the live feed", "9,262", "`control-service/src/seed/harvest.ts` (2026-08-09)"],
        ["Records with `status == 'Live'`", "664", "`harvest.ts`"],
        ["Records carrying a route name", "1,602", "`harvest.ts`"],
        ["Distinct route names among the live subset", "517", "`harvest.ts`"],
        ["Live payload size", "~11.7 MB", "`scheduler/gpsPoll.ts`"],
        ["Active route-directions seeded", "666", "`control-service/src/seed/odTimetable.ts` (2026-08-11)"],
        ["Route-directions with a target headway", "171  (74 timetable + 97 OD)", "`odTimetable.ts`"],
        ["Route-directions with none", "495", "`odTimetable.ts`"],
        ["OD corpus swept", "53,898 rows, 210 city pairs", "`odTimetable.ts`"],
        ["Median H*, bucketed estimator", "2,582 s  (~43 min)", "`odTimetable.ts`"],
        ["Median H*, pooled estimator (rejected)", "1,635 s", "`odTimetable.ts`"],
        ["Average route-direction length", "215 km, 19 stops", "`control-service/src/rehearsal/run.ts`"],
        ["Median inter-stop gap", "7.9 km", "`control-service/src/arrival-prediction/predict.ts`"],
        ["Median running speed, moving fleet", "38.5 km/h  (512 live rows, 2026-08-14)", "`predict.ts`"],
        ["Default max hold", "90 s", "`control-service/db/migrations/…__core_data_model.sql`"],
        ["Default bunched / warning thresholds", "0.25 × H*  /  0.5 × H*", "same migration"],
        ["Hard staleness bound on a candidate", "90 s", "`control-service/src/mpc/safety.ts`"],
        ["GPS poll interval / headway sweep interval", "30 s  /  60 s", "`control-service/src/config/env.ts`"],
        ["Headway batch size / concurrency / pool max", "60  /  4  /  10", "`config/env.ts`, `db/pool.ts`"],
        ["Fleet map render cost", "9,170 vehicles at 4.3 ms/frame", "`src/components/map/fleetCanvasLayer.ts`"],
        ["Fleet map cost before the rewrite", "53.7 ms/frame, 18.5 fps", "`fleetCanvasLayer.ts`"],
        ["Tests passing", "2,410  (731 + 1,679)", "Run for this evaluation"],
    ],
    widths=[2.5, 1.9, 2.6], fs=8.2)

h2("Projections, clearly separated")
table(
    ["Projection", "Basis", "Result"],
    [["Ingestion cycle cost at 14,000 buses",
      "fleet × 3 queries × round-trip, from the sequential loop in `ingestion/pipeline.ts` and the three queries "
      "in `state-estimation/service.ts`",
      "≈84 s at 2 ms RTT; ≈210 s at 5 ms"],
     ["Fleet size at which the cycle exceeds 30 s",
      "same model",
      "5,000 buses at 2 ms; 2,000 at 5 ms"],
     ["Detection latency at fleet scale",
      "3 samples × 60 s × ⌈eligible ÷ 60⌉, with eligible route-directions estimated at 400–600",
      "20–30 minutes"]],
    widths=[1.9, 3.3, 1.8], fs=8.2)

h1("Appendix B — How This Evaluation Was Made")
bullet([("Scope.", "b"), ("  All 98 commits between `main` and `integration/main`, read as a diff and as a "
         "final state.", "")])
bullet([("Verification.", "b"), ("  Both packages installed from the lockfile; test suites, typecheck and lint "
         "run in each. Results in §3 are from those runs.", "")])
bullet([("Sources.", "b"), ("  Quantities were taken from the repository's own measurement comments, which "
         "record the date and method of each measurement, and cross-checked against the code they describe.", "")])
bullet([("Limits.", "b"), ("  No live UPSRTC endpoint was contacted and no credentials were used. The E2E suite "
         "was not run — it needs services CI provisions. Throughput figures are arithmetic, not benchmarked; "
         "producing a real benchmark is recommendation #2.", "")])

doc.save(OUT)
print("wrote", OUT)
